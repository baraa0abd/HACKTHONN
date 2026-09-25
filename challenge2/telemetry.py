"""Parse and analyse the public BIRDS on-orbit EPS telemetry workbooks."""
from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import hashlib
import math
import statistics

from openpyxl import load_workbook


FACES = ("py", "px", "mz", "mx", "pz")
REQUIRED = ["Time Stamp", "Vbat (V)", "Ibatt(mA)"] + [f"V{x} (mV)" for x in FACES] + [f"I{x} (mA)" for x in FACES]


@dataclass(frozen=True)
class Sample:
    time_s: float
    solar_w: float
    battery_v: float
    battery_current_a: float
    battery_power_w: float
    face_power_w: dict[str, float]


def _number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    value = float(value)
    return value if math.isfinite(value) else None


def load_passes(data_dir: Path) -> tuple[list[dict], dict]:
    """Load operational sheets. Rows with incomplete calibrated channels are reported and skipped."""
    passes = []
    provenance = {}
    for book_path in sorted(data_dir.glob("*.xlsx")):
        satellite = book_path.stem
        provenance[book_path.name] = hashlib.sha256(book_path.read_bytes()).hexdigest()
        book = load_workbook(book_path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            if sheet.title.lower().startswith("test"):
                continue
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            normalized_header = {str(value).strip().casefold(): index for index, value in enumerate(header)}
            columns = {name: normalized_header[name.strip().casefold()] for name in REQUIRED
                       if name.strip().casefold() in normalized_header}
            missing = [name for name in REQUIRED if name not in columns]
            if missing:
                book.close()
                raise ValueError(f"{book_path.name}/{sheet.title}: missing columns {missing}")
            samples, rejected = [], 0
            for row in rows:
                values = {name: _number(row[index]) for name, index in columns.items()}
                if any(value is None for value in values.values()):
                    rejected += 1
                    continue
                faces = {face: max(0.0, values[f"V{face} (mV)"] * values[f"I{face} (mA)"] / 1_000_000)
                         for face in FACES}
                battery_current_a = values["Ibatt(mA)"] / 1000
                samples.append(Sample(values["Time Stamp"], sum(faces.values()), values["Vbat (V)"],
                                      battery_current_a, values["Vbat (V)"] * battery_current_a, faces))
            if len(samples) < 2:
                book.close()
                raise ValueError(f"{book_path.name}/{sheet.title}: fewer than two valid rows")
            if any(b.time_s <= a.time_s for a, b in zip(samples, samples[1:])):
                book.close()
                raise ValueError(f"{book_path.name}/{sheet.title}: timestamps are not strictly increasing")
            passes.append({"satellite": satellite, "pass": sheet.title, "samples": samples,
                           "rows_accepted": len(samples), "rows_rejected": rejected})
        book.close()
    if not passes:
        raise ValueError(f"No operational XLSX telemetry found in {data_dir}")
    return passes, provenance


def _integrate(samples: list[Sample], attribute: str) -> float:
    watt_seconds = 0.0
    for left, right in zip(samples, samples[1:]):
        watt_seconds += (getattr(left, attribute) + getattr(right, attribute)) * 0.5 * (right.time_s - left.time_s)
    return watt_seconds / 3600


def _percentile(values, fraction):
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    low = math.floor(position)
    high = math.ceil(position)
    return ordered[low] if low == high else ordered[low] * (high - position) + ordered[high] * (position - low)


def analyse_pass(record: dict, sun_threshold_w: float = 0.05) -> dict:
    samples = record["samples"]
    solar = [s.solar_w for s in samples]
    battery_v = [s.battery_v for s in samples]
    duration = samples[-1].time_s - samples[0].time_s
    dt = [b.time_s - a.time_s for a, b in zip(samples, samples[1:])]
    face_energy = {}
    face_zero_fraction = {}
    for face in FACES:
        energy = 0.0
        for left, right in zip(samples, samples[1:]):
            energy += (left.face_power_w[face] + right.face_power_w[face]) * 0.5 * (right.time_s-left.time_s) / 3600
        face_energy[face] = energy
        face_zero_fraction[face] = sum(s.face_power_w[face] <= .001 for s in samples) / len(samples)
    active_face_energies = [v for v in face_energy.values() if v > 0]
    peer_median = statistics.median(active_face_energies) if active_face_energies else 0
    flags = [f"{face.upper()} panel produced <1 mW for >90% of samples"
             for face in FACES if face_zero_fraction[face] > .9 and peer_median > .005]
    if min(battery_v) < 3.0 or max(battery_v) > 4.5:
        flags.append("battery voltage left the documented 3.0–4.5 V analysis guard band")
    return {"satellite": record["satellite"], "pass": record["pass"],
            "samples": len(samples), "rows_rejected": record["rows_rejected"],
            "duration_min": duration / 60, "median_sample_interval_s": statistics.median(dt),
            "solar_mean_w": statistics.fmean(solar), "solar_p95_w": _percentile(solar, .95),
            "solar_max_w": max(solar), "solar_energy_wh": _integrate(samples, "solar_w"),
            "sunlit_fraction_threshold": sum(p > sun_threshold_w for p in solar) / len(solar),
            "sun_threshold_w": sun_threshold_w,
            "battery_v_min": min(battery_v), "battery_v_median": statistics.median(battery_v),
            "battery_v_max": max(battery_v),
            "battery_net_terminal_wh": _integrate(samples, "battery_power_w"),
            "battery_charge_fraction": sum(s.battery_current_a < 0 for s in samples) / len(samples),
            "face_energy_wh": face_energy, "face_zero_fraction": face_zero_fraction,
            "health_flags": flags}


def analyse_dataset(data_dir: Path) -> dict:
    passes, provenance = load_passes(data_dir)
    results = [analyse_pass(record) for record in passes]
    satellite = {}
    for name in sorted({r["satellite"] for r in results}):
        group = [r for r in results if r["satellite"] == name]
        satellite[name] = {"passes": len(group), "samples": sum(r["samples"] for r in group),
                           "rejected_rows": sum(r["rows_rejected"] for r in group),
                           "median_solar_energy_wh": statistics.median(r["solar_energy_wh"] for r in group),
                           "median_solar_mean_w": statistics.median(r["solar_mean_w"] for r in group),
                           "min_battery_v": min(r["battery_v_min"] for r in group),
                           "max_battery_v": max(r["battery_v_max"] for r in group),
                           "flagged_passes": sum(bool(r["health_flags"]) for r in group)}
    return {"method": {"panel_power": "sum(max(0, V_mV × I_mA / 1,000,000)) across five faces",
                       "solar_energy": "trapezoidal integration over recorded timestamps",
                       "battery_sign": "positive terminal Wh is discharge; negative is charge",
                       "sunlit_proxy": "aggregate measured panel power > 0.05 W",
                       "invalid_rows": "skip and count rows missing any required calibrated numeric channel"},
            "provenance_sha256": provenance, "satellites": satellite, "passes": results}
